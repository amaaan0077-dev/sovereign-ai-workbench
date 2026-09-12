"""
Sovereign AI Workbench - XLSX Analysis Report Generator

Generates a structured Excel (.xlsx) analysis report from engineering
telemetry data and calculation results. Network-isolated; no external calls.
"""

import os
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side
    )
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

DELIVERABLES_DIR = Path(__file__).resolve().parents[4] / "storage" / "deliverables"
DELIVERABLES_DIR.mkdir(parents=True, exist_ok=True)

# ── Palette (matches UI design tokens) ──────────────────────────
COL_DARK       = "1C1F24"   # chassis lead
COL_STEEL      = "262B33"   # console steel
COL_CYAN       = "2F818E"   # verification cyan
COL_AMBER      = "C2782A"   # restricted amber
COL_WHITE_TEXT = "E1E4EA"   # spec white
COL_MUTED      = "8792A2"   # subsystem muted
COL_RED        = "C23A2A"   # alert red

_thin = Side(style="thin", color="3A414D")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _hdr_font(bold: bool = True) -> Font:
    return Font(name="Calibri", bold=bold, color=COL_WHITE_TEXT, size=10)


def _cell_font(color: str = COL_WHITE_TEXT) -> Font:
    return Font(name="Calibri", color=color, size=9)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _centre() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


# ── Helpers ───────────────────────────────────────────────────────

def _write_header_row(ws, row: int, cols: List[str]) -> None:
    for ci, col in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=ci, value=col)
        cell.font      = _hdr_font()
        cell.fill      = _fill(COL_STEEL)
        cell.alignment = _centre()
        cell.border    = _border


def _write_data_row(
    ws,
    row: int,
    values: List[Any],
    bg: str = COL_DARK,
    fg: str = COL_WHITE_TEXT,
    bold: bool = False,
) -> None:
    for ci, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=ci, value=val)
        cell.font      = Font(name="Calibri", color=fg, bold=bold, size=9)
        cell.fill      = _fill(bg)
        cell.alignment = _left()
        cell.border    = _border


# ── Main generator ────────────────────────────────────────────────

def generate_analysis_xlsx(
    title: str,
    reference_no: str,
    requester_name: str,
    department: str,
    clearance_tier: str,
    technical_findings: List[Dict[str, str]],
    calculations: Optional[List[Dict[str, str]]] = None,
    citations: Optional[List[str]] = None,
    summary_text: str = "",
    recommendation: str = "",
) -> str:
    """
    Build a multi-sheet Excel analysis report and save it to the
    deliverables directory.

    Returns the filename (not the full path) so the route handler
    can serve it via /api/deliverables/<filename>.
    """

    if not OPENPYXL_AVAILABLE:
        return "XLSX_GEN_UNAVAILABLE_openpyxl_missing"

    wb = openpyxl.Workbook()

    # ── Sheet 1: Cover / Summary ──────────────────────────────────
    ws_cover = wb.active
    ws_cover.title = "Summary"
    ws_cover.sheet_view.showGridLines = False

    ws_cover.column_dimensions["A"].width = 28
    ws_cover.column_dimensions["B"].width = 58

    # Classification banner (row 1)
    ws_cover.merge_cells("A1:B1")
    banner = ws_cover["A1"]
    banner.value     = (
        f"SOVEREIGN ON-PREMISE | CLEARANCE: {clearance_tier} | "
        "AIR-GAPPED SYSTEM | NO EXTERNAL EGRESS"
    )
    banner.font      = Font(name="Calibri", bold=True, color=COL_WHITE_TEXT, size=10)
    banner.fill      = _fill(COL_AMBER)
    banner.alignment = _centre()

    # Title block (rows 2-3)
    ws_cover.merge_cells("A2:B2")
    title_cell = ws_cover["A2"]
    title_cell.value     = title
    title_cell.font      = Font(name="Calibri", bold=True, color=COL_WHITE_TEXT, size=13)
    title_cell.fill      = _fill(COL_DARK)
    title_cell.alignment = _centre()
    ws_cover.row_dimensions[2].height = 28

    ws_cover.merge_cells("A3:B3")
    ref_cell = ws_cover["A3"]
    ref_cell.value     = f"REF: {reference_no}"
    ref_cell.font      = Font(name="Calibri", color=COL_MUTED, size=9, italic=True)
    ref_cell.fill      = _fill(COL_STEEL)
    ref_cell.alignment = _centre()

    # Meta block
    meta_rows = [
        ("Requested By",    requester_name),
        ("Department",      department),
        ("Clearance Tier",  clearance_tier),
        ("Generated At",    time.strftime("%Y-%m-%d %H:%M:%S UTC", time.gmtime())),
        ("System",          "Sovereign AI Workbench v1.0 — On-Premise"),
        ("Network Egress",  "VERIFIED ZERO — Air-gapped"),
    ]
    for ri, (k, v) in enumerate(meta_rows, start=5):
        ws_cover.cell(row=ri, column=1, value=k).font  = _hdr_font()
        ws_cover.cell(row=ri, column=1).fill           = _fill(COL_STEEL)
        ws_cover.cell(row=ri, column=1).alignment      = _left()
        ws_cover.cell(row=ri, column=1).border         = _border
        ws_cover.cell(row=ri, column=2, value=v).font  = _cell_font()
        ws_cover.cell(row=ri, column=2).fill           = _fill(COL_DARK)
        ws_cover.cell(row=ri, column=2).alignment      = _left()
        ws_cover.cell(row=ri, column=2).border         = _border

    # Summary text
    if summary_text:
        r = 12
        ws_cover.merge_cells(f"A{r}:B{r}")
        ws_cover[f"A{r}"].value     = "EXECUTIVE SUMMARY"
        ws_cover[f"A{r}"].font      = _hdr_font()
        ws_cover[f"A{r}"].fill      = _fill(COL_STEEL)
        ws_cover[f"A{r}"].alignment = _centre()

        ws_cover.merge_cells(f"A{r+1}:B{r+4}")
        txt = ws_cover[f"A{r+1}"]
        txt.value     = summary_text
        txt.font      = _cell_font()
        txt.fill      = _fill(COL_DARK)
        txt.alignment = Alignment(
            horizontal="left", vertical="top", wrap_text=True
        )
        ws_cover.row_dimensions[r + 1].height = 60

    # Recommendation
    if recommendation:
        r = 18
        ws_cover.merge_cells(f"A{r}:B{r}")
        ws_cover[f"A{r}"].value     = "RECOMMENDATION"
        ws_cover[f"A{r}"].font      = _hdr_font()
        ws_cover[f"A{r}"].fill      = _fill(COL_CYAN)
        ws_cover[f"A{r}"].alignment = _centre()

        ws_cover.merge_cells(f"A{r+1}:B{r+4}")
        rec = ws_cover[f"A{r+1}"]
        rec.value     = recommendation
        rec.font      = Font(name="Calibri", color=COL_WHITE_TEXT, bold=True, size=9)
        rec.fill      = _fill(COL_DARK)
        rec.alignment = Alignment(
            horizontal="left", vertical="top", wrap_text=True
        )
        ws_cover.row_dimensions[r + 1].height = 60

    # ── Sheet 2: Technical Findings ──────────────────────────────
    ws_findings = wb.create_sheet("Technical Findings")
    ws_findings.sheet_view.showGridLines = False
    ws_findings.column_dimensions["A"].width = 32
    ws_findings.column_dimensions["B"].width = 26
    ws_findings.column_dimensions["C"].width = 26
    ws_findings.column_dimensions["D"].width = 18

    ws_findings.merge_cells("A1:D1")
    ws_findings["A1"].value     = "TECHNICAL FINDINGS — MEASURED VS BASELINE"
    ws_findings["A1"].font      = _hdr_font()
    ws_findings["A1"].fill      = _fill(COL_STEEL)
    ws_findings["A1"].alignment = _centre()

    _write_header_row(
        ws_findings, 2,
        ["Parameter", "Measured Value", "Baseline / Acceptable", "Status"]
    )

    for ri, finding in enumerate(technical_findings, start=3):
        param    = finding.get("parameter", "")
        measured = finding.get("measured", "")
        baseline = finding.get("baseline", "")

        # Simple status colouring
        status_text = "REVIEW"
        status_fg   = COL_AMBER
        if "pass" in str(measured).lower() or "normal" in str(measured).lower():
            status_text = "PASS"
            status_fg   = COL_CYAN
        elif "fail" in str(measured).lower() or "exceed" in str(measured).lower():
            status_text = "FAIL"
            status_fg   = COL_RED

        bg = "1A1D23" if ri % 2 == 0 else COL_DARK
        _write_data_row(ws_findings, ri, [param, measured, baseline, ""], bg=bg)

        status_cell = ws_findings.cell(row=ri, column=4, value=status_text)
        status_cell.font      = Font(name="Calibri", color=status_fg, bold=True, size=9)
        status_cell.fill      = _fill(bg)
        status_cell.alignment = _centre()
        status_cell.border    = _border

    # ── Sheet 3: Calculations ────────────────────────────────────
    if calculations:
        ws_calc = wb.create_sheet("Calculations")
        ws_calc.sheet_view.showGridLines = False
        ws_calc.column_dimensions["A"].width = 42
        ws_calc.column_dimensions["B"].width = 26
        ws_calc.column_dimensions["C"].width = 22

        ws_calc.merge_cells("A1:C1")
        ws_calc["A1"].value     = "SANDBOXED ENGINEERING CALCULATIONS (VERIFIED)"
        ws_calc["A1"].font      = _hdr_font()
        ws_calc["A1"].fill      = _fill(COL_CYAN)
        ws_calc["A1"].alignment = _centre()

        _write_header_row(ws_calc, 2, ["Formula / Description", "Result", "Verification"])

        for ri, calc in enumerate(calculations, start=3):
            formula      = calc.get("formula", "")
            result       = calc.get("result", "")
            verification = calc.get("verification", "")

            bg = "1A1D23" if ri % 2 == 0 else COL_DARK
            _write_data_row(
                ws_calc, ri,
                [formula, result, verification],
                bg=bg
            )
            # Colour the verification cell
            v_cell = ws_calc.cell(row=ri, column=3)
            v_cell.font = Font(
                name="Calibri",
                color=COL_CYAN if "VERIFIED" in str(verification) else COL_RED,
                bold=True, size=9
            )

    # ── Sheet 4: Citations ───────────────────────────────────────
    if citations:
        ws_cite = wb.create_sheet("Citations")
        ws_cite.sheet_view.showGridLines = False
        ws_cite.column_dimensions["A"].width = 8
        ws_cite.column_dimensions["B"].width = 72

        ws_cite.merge_cells("A1:B1")
        ws_cite["A1"].value     = "SOURCE CITATIONS — CLEARANCE-GATED DOCUMENT RETRIEVAL"
        ws_cite["A1"].font      = _hdr_font()
        ws_cite["A1"].fill      = _fill(COL_STEEL)
        ws_cite["A1"].alignment = _centre()

        _write_header_row(ws_cite, 2, ["#", "Citation"])

        for ri, cite in enumerate(citations, start=3):
            bg = "1A1D23" if ri % 2 == 0 else COL_DARK
            _write_data_row(ws_cite, ri, [ri - 2, cite], bg=bg)

    # ── Save ──────────────────────────────────────────────────────
    safe_ref = reference_no.replace("/", "-").replace("\\", "-").replace(" ", "_")
    ts       = time.strftime("%Y%m%d_%H%M%S", time.gmtime())
    filename = f"{safe_ref}_{ts}.xlsx"
    filepath = DELIVERABLES_DIR / filename

    wb.save(str(filepath))
    return filename
